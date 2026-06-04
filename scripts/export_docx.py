#!/usr/bin/env python3
"""Export NKP reference markdown to one .docx file per version."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", default=".", help="Repo root containing version folders.")
    parser.add_argument("--version", default="", help="Optional single version to export.")
    parser.add_argument(
        "--output-dir",
        default="exports",
        help="Directory to write generated .docx files.",
    )
    parser.add_argument(
        "--template",
        default="",
        help="Optional path to a .docx template file for styles/layout.",
    )
    parser.add_argument(
        "--replace-template-content",
        action="store_true",
        help="Deprecated: template content is cleared by default.",
    )
    parser.add_argument(
        "--keep-template-content",
        action="store_true",
        help="Keep existing template body content and append generated content after it.",
    )
    return parser.parse_args()


def discover_versions(root: Path, selected: str) -> list[str]:
    if selected:
        version_dir = root / selected
        if not version_dir.is_dir():
            raise RuntimeError(f"Version directory not found: {version_dir}")
        return [selected]
    return sorted(
        p.name for p in root.iterdir() if p.is_dir() and re.match(r"^v?\d+\.\d+\.\d+$", p.name)
    )


def all_pages(version_dir: Path) -> list[Path]:
    pages: list[Path] = []
    root_page = version_dir / "nkp.md"
    if root_page.exists():
        pages.append(root_page)

    commands_root = version_dir / "commands"
    if not commands_root.is_dir():
        return pages

    def walk_command_dir(directory: Path) -> None:
        # Emit command page(s) in this directory first.
        # File names are full command slugs, e.g. nkp-create-cluster.md.
        dir_pages = sorted(directory.glob("nkp-*.md"))
        pages.extend(dir_pages)

        # Emit child command directories in name order.
        child_dirs = sorted([p for p in directory.iterdir() if p.is_dir()])
        for child in child_dirs:
            walk_command_dir(child)

    top_level_dirs = sorted([p for p in commands_root.iterdir() if p.is_dir()])
    for cmd_dir in top_level_dirs:
        walk_command_dir(cmd_dir)

    return pages


def normalize_text(text: str) -> str:
    text = text.replace("`", "")
    # Convert markdown links to visible labels.
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text


def add_bookmark(paragraph, bookmark_id: int, name: str) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_hyperlink(paragraph, text: str, anchor: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    run.append(text_elem)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def normalize_heading_paragraph(paragraph) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    style_names = {s.name for s in doc.styles}
    if "Table Grid" in style_names:
        table.style = "Table Grid"

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.append(tbl_pr)
    tbl_borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "8")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), "000000")
        tbl_borders.append(elem)
    tbl_pr.append(tbl_borders)

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), "D9D9D9")
        tc_pr.append(shading)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0, 0, 0)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value


def render_help_block(doc: Document, lines: list[str]) -> None:
    usage_mode = False
    example_mode = False
    continuation_mode = False
    i = 0
    total = len(lines)

    while i < total:
        line = lines[i]
        stripped = line.strip()

        if stripped == "":
            doc.add_paragraph("")
            usage_mode = False
            example_mode = False
            continuation_mode = False
            i += 1
            continue

        heading_match = re.match(r"^\s*(#{1,6})\s+(.*)$", line)
        if heading_match:
            level = len(heading_match.group(1))
            text = normalize_text(heading_match.group(2).strip())
            if level >= 4:
                # Keep deep "heading-like" labels (e.g. Linux/macOS) out of doc navigation.
                p = doc.add_paragraph()
                p.add_run(text).bold = True
            else:
                p = doc.add_heading(text, level=min(level, 6))
                normalize_heading_paragraph(p)
            i += 1
            continue

        if stripped == "Available Commands:":
            i += 1
            rows: list[list[str]] = []
            while i < total:
                row_line = lines[i]
                if row_line.strip() == "":
                    i += 1
                    break
                m = re.match(r"^\s{2,}(\S+)\s{2,}(.*)$", row_line)
                if not m:
                    break
                rows.append([normalize_text(m.group(1)), normalize_text(m.group(2).strip())])
                i += 1
            if rows:
                p = doc.add_paragraph("Available Commands")
                p.runs[0].bold = True
                p.paragraph_format.space_before = Pt(10)
                add_table(doc, ["Command", "Description"], rows)
            continue

        if stripped in {"Flags:", "Global Flags:", "Inherited Flags:"}:
            section_title = stripped.rstrip(":")
            i += 1
            rows = []
            while i < total:
                row_line = lines[i]
                if row_line.strip() == "":
                    i += 1
                    break
                m = re.match(r"^\s{2,}(.+?)\s{2,}(.*)$", row_line)
                if not m:
                    break
                rows.append([normalize_text(m.group(1).strip()), normalize_text(m.group(2).strip())])
                i += 1
            if rows:
                p = doc.add_paragraph(section_title)
                p.runs[0].bold = True
                p.paragraph_format.space_before = Pt(10)
                add_table(doc, ["Flag", "Description"], rows)
            continue

        if stripped.startswith("Usage:"):
            usage_mode = True
            doc.add_paragraph(normalize_text(line))
            i += 1
            continue

        if stripped.startswith("Examples:") or stripped.startswith("Example:"):
            example_mode = True
            doc.add_paragraph(normalize_text(line))
            i += 1
            continue

        looks_like_command = bool(
            re.match(
                r'^\s*(\$|nkp\b|kubectl\b|docker\b|helm\b|az\b|aws\b|gcloud\b|source\b|echo\b)',
                line,
            )
        )
        style_as_code = looks_like_command or (usage_mode and line.startswith("  ")) or (
            example_mode and line.startswith("  ")
        )
        if continuation_mode and stripped != "":
            style_as_code = True
        p = doc.add_paragraph(normalize_text(line))
        if style_as_code:
            for run in p.runs:
                run.font.name = "Courier New"
                run.font.size = Pt(9)
        continuation_mode = style_as_code and stripped.endswith("\\")
        i += 1


def pick_style(doc: Document, style_names: list[str]) -> str | None:
    existing = {s.name for s in doc.styles}
    for name in style_names:
        if name in existing:
            return name
    return None


def toc_style_for_depth(doc: Document, depth: int) -> str | None:
    level = min(depth + 1, 9)
    candidates = [f"TOC {level}", f"TOC{level}", f"Contents {level}", f"Contents{level}"]
    return pick_style(doc, candidates)


def render_markdown(doc: Document, md_text: str, skip_first_h1: bool = False) -> None:
    in_code = False
    code_lines: list[str] = []
    skipped_h1 = False
    heading_as_bold_text = {"subcommands", "inlined subcommands"}

    style_names = {s.name for s in doc.styles}
    has_list_bullet = "List Bullet" in style_names

    for raw_line in md_text.splitlines():
        line = raw_line.rstrip("\n")
        normalized_line = normalize_text(line).strip()

        if normalized_line == "Inlined Subcommands":
            continue
        if normalized_line == "_These are kept on this page to avoid excessive page-per-provider splits._":
            continue

        if line.strip().startswith("```"):
            if in_code:
                render_help_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading_match:
            level = len(heading_match.group(1))
            text = normalize_text(heading_match.group(2).strip())
            if text.lower() == "inlined subcommands":
                continue
            if level == 1 and skip_first_h1 and not skipped_h1:
                skipped_h1 = True
                continue
            if text.lower() in heading_as_bold_text:
                p = doc.add_paragraph()
                p.add_run(text).bold = True
            else:
                p = doc.add_heading(text, level=min(level, 6))
                normalize_heading_paragraph(p)
            continue
        if line.startswith("- "):
            bullet_text = normalize_text(line[2:].strip())
            if has_list_bullet:
                doc.add_paragraph(bullet_text, style="List Bullet")
            else:
                p = doc.add_paragraph()
                p.add_run("• ")
                p.add_run(bullet_text)
            continue
        if line.strip() == "":
            doc.add_paragraph("")
            continue

        doc.add_paragraph(normalize_text(line))

    if in_code and code_lines:
        render_help_block(doc, code_lines)


def heading_for_page(page: Path) -> str:
    text = page.read_text(encoding="utf-8")
    for line in text.splitlines():
        if line.startswith("# "):
            return normalize_text(line[2:].strip())
    return normalize_text(page.stem.replace("nkp-", "", 1))


def nav_label(version_dir: Path, page: Path) -> tuple[str, int]:
    if page.name == "nkp.md":
        return "nkp", 0
    rel = page.relative_to(version_dir / "commands")
    depth = max(0, len(rel.parts) - 2)
    command = page.stem
    return command, depth


def clear_document_content(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def trim_to_cover_page(doc: Document) -> None:
    """Keep template content up to the first explicit page break."""
    body = doc._body._element
    keep_count = 0
    page_break_found = False

    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        keep_count += 1
        # Keep content through first explicit page break marker.
        if child.xpath('.//*[local-name()="br" and @*[local-name()="type"]="page"]'):
            page_break_found = True
            break
        if child.xpath('.//*[local-name()="lastRenderedPageBreak"]'):
            page_break_found = True
            break

    children = [c for c in list(body) if not c.tag.endswith("sectPr")]
    if not page_break_found:
        # If no page break was found, keep only the first body element.
        keep_count = 1 if children else 0

    for idx, child in enumerate(children):
        if idx >= keep_count:
            body.remove(child)


def build_version_doc(
    root: Path,
    output_dir: Path,
    version: str,
    template_path: Path | None,
    keep_template_content: bool,
) -> Path:
    version_dir = root / version
    out_path = output_dir / f"NKP_CLI_Reference_{version}.docx"

    doc = Document(str(template_path)) if template_path else Document()
    if template_path and not keep_template_content:
        clear_document_content(doc)
    elif template_path and keep_template_content:
        doc.add_page_break()

    title = doc.add_heading(f"NKP CLI Reference {version}", level=0)
    normalize_heading_paragraph(title)
    pages = all_pages(version_dir)
    bookmarks: dict[Path, str] = {}
    for idx, page in enumerate(pages, start=1):
        bookmarks[page] = f"sec_{idx}"

    for page in pages:
        label, depth = nav_label(version_dir, page)
        toc_style = toc_style_for_depth(doc, depth)
        nav_par = doc.add_paragraph(style=toc_style) if toc_style else doc.add_paragraph()
        if not toc_style:
            nav_par.paragraph_format.left_indent = Pt(0)
            nav_par.paragraph_format.first_line_indent = Pt(0)
            prefix = ("  " * depth) + "• "
            nav_par.add_run(prefix)
        add_internal_hyperlink(nav_par, normalize_text(label), bookmarks[page])

    for idx, page in enumerate(pages, start=1):
        doc.add_page_break()
        _, depth = nav_label(version_dir, page)
        heading_level = min(depth + 1, 6)
        heading = doc.add_heading(heading_for_page(page), level=heading_level)
        normalize_heading_paragraph(heading)
        add_bookmark(heading, idx, bookmarks[page])
        render_markdown(doc, page.read_text(encoding="utf-8"), skip_first_h1=True)

    output_dir.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def main() -> None:
    args = parse_args()
    root = Path(args.root).resolve()
    output_dir = Path(args.output_dir).resolve()
    template_path: Path | None = None
    if args.template:
        candidate = Path(args.template)
        if not candidate.is_absolute():
            candidate = (root / candidate).resolve()
        if not candidate.exists():
            raise RuntimeError(f"Template file not found: {candidate}")
        template_path = candidate

    versions = discover_versions(root, args.version)
    if not versions:
        raise RuntimeError(f"No version folders found in: {root}")

    for version in versions:
        out_path = build_version_doc(
            root,
            output_dir,
            version,
            template_path,
            args.keep_template_content,
        )
        print(f"Wrote: {out_path}")


if __name__ == "__main__":
    main()
